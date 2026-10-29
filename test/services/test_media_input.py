import asyncio
import base64
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from PIL import Image

from services.llm_client import LLMClient
from services.media_input import MAX_DOCUMENT_CHARS, MediaInputReader


def image_bytes():
    buffer = BytesIO()
    Image.new("RGB", (120, 80), "red").save(buffer, format="PNG")
    return buffer.getvalue()


def user_message(filename, *, image=False):
    return {"role": "user", "content": "请分析上传内容", "images" if image else "attachments": [
        {"name": filename, "url": f"/generated/uploads/{filename}", "path": f"uploads/{filename}"}
    ]}


@pytest.fixture
def media_dir(tmp_path):
    (tmp_path / "uploads").mkdir()
    return tmp_path


def test_images_are_sent_as_valid_base64_and_keep_user_text(media_dir):
    (media_dir / "uploads/photo.png").write_bytes(image_bytes())
    message = user_message("photo.png", image=True)
    content = LLMClient._chat_messages([message], media_dir)[0]["content"]
    assert content[0]["type"] == "text"
    assert "请分析上传内容" in content[0]["text"]
    url = content[1]["image_url"]["url"]
    with Image.open(BytesIO(base64.b64decode(url.split(",")[1]))) as image:
        assert image.size == (120, 80)
        assert image.getpixel((0, 0))[0] > 200
    assert message["images"][0]["url"].startswith("/generated/")
    assert "data:image" not in str(message)  # Never persist Base64 in session metadata.


def test_transparent_images_keep_a_readable_white_background(media_dir):
    image = Image.new("RGBA", (4000, 2000), (0, 0, 0, 0))
    image.save(media_dir / "uploads/transparent.png")
    content = MediaInputReader(media_dir).content(user_message("transparent.png", image=True))
    url = content[1]["image_url"]["url"]
    with Image.open(BytesIO(base64.b64decode(url.split(",")[1]))) as decoded:
        assert decoded.size == (2048, 1024)
        assert decoded.getpixel((0, 0)) == (255, 255, 255)


@pytest.mark.parametrize("suffix,body", [("txt", "发票金额 123.45"), ("csv", "姓名,金额\n张三,15"), ("json", '{"total": 42}')])
def test_text_documents_reach_model(media_dir, suffix, body):
    (media_dir / f"uploads/data.{suffix}").write_text(body, encoding="utf-8")
    content = LLMClient._chat_messages([user_message(f"data.{suffix}")], media_dir)[0]["content"]
    assert body in content
    assert "<document>" in content


def test_docx_paragraphs_and_tables(media_dir):
    from docx import Document
    doc = Document()
    doc.add_paragraph("合同金额 500 元")
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = "乙方：测试公司"
    doc.save(media_dir / "uploads/contract.docx")
    text, _ = MediaInputReader(media_dir).document({"path": "uploads/contract.docx"})
    assert "合同金额 500 元" in text
    assert "乙方：测试公司" in text


def test_xlsx_sheets_values_and_formulas(media_dir):
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "销售"
    book.active.append(["商品", "收入"])
    book.active.append(["茶叶", 125])
    book.create_sheet("汇总").append(["=SUM(销售!B2:B2)"])
    book.save(media_dir / "uploads/sales.xlsx")
    text, _ = MediaInputReader(media_dir).document({"path": "uploads/sales.xlsx"})
    assert "工作表: 销售" in text and "茶叶\t125" in text
    assert "工作表: 汇总" in text and "=SUM(" in text


def test_pptx_text_and_tables(media_dir):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "季度报告"
    shape = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    shape.table.cell(0, 1).text = "收入 900 元"
    deck.save(media_dir / "uploads/report.pptx")
    text, _ = MediaInputReader(media_dir).document({"path": "uploads/report.pptx"})
    assert "季度报告" in text and "收入 900 元" in text and "幻灯片 1" in text


def test_pdf_text_and_scanned_pages_reach_vision_model(media_dir):
    import pymupdf
    with pymupdf.open() as doc:
        doc.new_page().insert_text((40, 60), "Invoice total: 123.45 EUR, paid in full.")
        doc.new_page().insert_image(pymupdf.Rect(20, 20, 140, 100), stream=image_bytes())
        doc.save(media_dir / "uploads/invoice.pdf")
    content = LLMClient._chat_messages([user_message("invoice.pdf")], media_dir)[0]["content"]
    assert "Invoice total: 123.45" in content[0]["text"]
    assert "第 2 页" in content[0]["text"]
    assert content[1]["type"] == "image_url"


def test_missing_corrupt_and_unsupported_files_are_explicit(media_dir):
    (media_dir / "uploads/broken.png").write_text("not a picture")
    (media_dir / "uploads/old.doc").write_bytes(b"old binary format")
    for name, image, expected in [
        ("gone.pdf", False, "重新上传"),
        ("broken.png", True, "UnidentifiedImageError"),
        ("old.doc", False, "请先转换"),
    ]:
        content = LLMClient._chat_messages([user_message(name, image=image)], media_dir)[0]["content"]
        assert isinstance(content, str) and "附件读取失败" in content and expected in content


def test_paths_cannot_escape_generated_directory_even_through_symlinks(media_dir, tmp_path):
    secret = tmp_path.parent / f"{tmp_path.name}-secret.txt"
    secret.write_text("SHOULD_NOT_BE_SENT")
    (media_dir / "uploads/link.txt").symlink_to(secret)
    reader = MediaInputReader(media_dir)
    for path in [str(secret), "uploads/link.txt", "/generated/%2e%2e/config/.env"]:
        with pytest.raises(ValueError, match="目录之外"):
            reader.local_path(path)
    with pytest.raises(ValueError, match="外部文件"):
        reader.local_path("https://example.com/private.txt")


def test_large_documents_are_bounded_and_cache_updates_after_file_changes(media_dir):
    path = media_dir / "uploads/large.txt"
    path.write_text("A" * (MAX_DOCUMENT_CHARS + 100))
    reader = MediaInputReader(media_dir)
    text, _ = reader.document({"path": "uploads/large.txt"})
    assert "已截断" in text and len(text) < MAX_DOCUMENT_CHARS + 100
    path.write_text("Updated content")
    assert reader.document({"path": "uploads/large.txt"})[0] == "Updated content"


def test_context_budget_and_summary_include_extracted_body(media_dir):
    from core.context_compressor import ContextCompressor
    (media_dir / "uploads/notes.txt").write_text("CONFIDENTIAL_DOCUMENT_BODY" * 100)
    compressor = ContextCompressor(SimpleNamespace(media_root=media_dir))
    message = user_message("notes.txt")
    assert compressor._count_chars([message]) > 2500
    assert compressor._count_tokens([message]) > 300
    assert "CONFIDENTIAL_DOCUMENT_BODY" in compressor._summary_content(message)


def test_every_llm_entry_point_uses_multimodal_payload(media_dir):
    (media_dir / "uploads/photo.png").write_bytes(image_bytes())
    messages = [user_message("photo.png", image=True)]
    llm = LLMClient("test-key", "https://dashscope.aliyuncs.com/compatible-mode/v1", "qwen3.8-27b", media_root=str(media_dir))
    # Close the unused real HTTP client before replacing it with a stub.
    llm.client.close()
    choice = SimpleNamespace(message=SimpleNamespace(content="a red picture", tool_calls=[]), finish_reason="stop")
    response = SimpleNamespace(choices=[choice])
    chunk = SimpleNamespace(choices=[SimpleNamespace(delta=SimpleNamespace(content="red"), finish_reason="stop")])
    create = Mock(return_value=response)
    llm.client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    assert llm.chat(messages) == "a red picture"
    llm.chat_with_tools(messages, [])
    create.return_value = [chunk]
    assert list(llm.stream_chat(messages)) == ["red"]

    async def run():
        async_create = AsyncMock(return_value=response)
        llm._async_client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=async_create)))
        assert await llm.achat(messages) == "a red picture"
        await llm.achat_with_tools(messages, [])
        async def chunks():
            yield chunk
        async_create.return_value = chunks()
        assert [part async for part in llm.astream_chat(messages)] == ["red"]
        async_create.return_value = chunks()
        events = [event async for event in llm.astream_with_tools(messages, [])]
        assert events[-1]["response"].content == "red"
        return async_create.call_args_list

    calls = create.call_args_list + asyncio.run(run())
    assert len(calls) == 7
    for call in calls:
        assert call.kwargs["model"] == "qwen3.8-27b"
        user = next(m for m in call.kwargs["messages"] if m["role"] == "user")
        assert user["content"][1]["type"] == "image_url"
    assert llm._request_payload([{"role": "user", "content": "hello"}])["model"] == "qwen3.8-27b"
    llm.vision_model = "custom-vision-model"
    assert llm._request_payload(messages)["model"] == "custom-vision-model"


def test_remote_images_are_forwarded_without_local_fetching():
    content = LLMClient._chat_messages([{"role": "user", "content": "看图", "images": ["https://example.com/photo.jpg"]}])[0]["content"]
    assert content[1]["image_url"]["url"] == "https://example.com/photo.jpg"


def test_conversion_failure_is_recorded_without_inline_image_data(media_dir):
    llm = LLMClient("test", "https://example.com/v1", "multimodal-model", media_root=str(media_dir))
    try:
        payload = llm._request_payload([user_message("missing.png", image=True)])
        assert payload["model"] == "multimodal-model"
        diagnostic = llm.last_request_diagnostics
        assert diagnostic["image_count"] == 0
        assert diagnostic["media_errors"][0]["error_type"] == "ValueError"
        assert "重新上传" in diagnostic["media_errors"][0]["reason"]
        assert "base64" not in str(diagnostic)
        llm._request_payload([{"role": "user", "content": "hello"}])
        assert llm.last_request_diagnostics["media_errors"] == []
    finally:
        llm.close()
