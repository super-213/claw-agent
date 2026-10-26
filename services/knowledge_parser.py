"""Persistent document extraction, independent of per-message truncation."""
from pathlib import Path
from zipfile import ZipFile

from services.media_input import IMAGE_EXTENSIONS, TEXT_EXTENSIONS, _image_data_url, _PDF_LOCK

MAX_PAGES = 1000
MAX_CHARS = 2_000_000
MAX_OCR_PAGES = 20


def extract_document(path, ocr=None):
    path = Path(path)
    units, warnings = [], []
    total, ocr_count = 0, 0

    def add(label, text, page=None):
        nonlocal total
        text = str(text or '').strip()
        remaining = MAX_CHARS - total
        if len(text) > remaining:
            text = text[:max(0, remaining)]
            warnings.append('已达到全文处理长度上限，部分内容未索引。')
        total += len(text)
        for start in range(0, len(text), 900):
            units.append({'location': label, 'page': page, 'text': text[start:start + 1100]})

    def recognize(data, label):
        nonlocal ocr_count
        if not ocr or ocr_count >= MAX_OCR_PAGES:
            warnings.append(f'{label}未识别，可重新处理或提供文字版。')
            return ''
        ocr_count += 1
        try:
            text = str(ocr(_image_data_url(data)) or '').strip()
            if not text:
                warnings.append(f'{label}未识别到文字。')
            return text
        except Exception:
            warnings.append(f'{label}文字识别失败，请重试。')
            return ''

    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import pymupdf
        with _PDF_LOCK, pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError('PDF 已加密，请先解除密码。')
            if len(doc) > MAX_PAGES:
                warnings.append(f'仅处理前 {MAX_PAGES} 页。')
            for index in range(min(len(doc), MAX_PAGES)):
                if total >= MAX_CHARS:
                    warnings.append('已达到全文处理长度上限，后续页面未索引。')
                    break
                page = doc[index]
                body = page.get_text()
                if len(body.strip()) < 20:
                    scale = min(1.5, 2048 / max(page.rect.width, page.rect.height))
                    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale))
                    body = recognize(pix.tobytes('png'), f'第 {index + 1} 页')
                add(f'第 {index + 1} 页', body, index + 1)
    elif suffix in IMAGE_EXTENSIONS:
        add('图片文字', recognize(path.read_bytes(), '图片'))
    elif suffix in TEXT_EXTENSIONS:
        data = path.read_bytes()
        for encoding in ('utf-8-sig', 'gb18030', 'utf-16'):
            try:
                body = data.decode(encoding)
                if '\x00' not in body:
                    break
            except UnicodeError:
                continue
        else:
            raise ValueError('无法识别文本编码。')
        add('正文', body)
    elif suffix in {'.docx', '.xlsx', '.pptx'}:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('解压后的文档过大。')
        if suffix == '.docx':
            from docx import Document
            doc = Document(path)
            for index, paragraph in enumerate(doc.paragraphs):
                add(f'段落 {index + 1}', paragraph.text)
            for index, table in enumerate(doc.tables):
                add(f'表格 {index + 1}', '\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows))
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        if total >= MAX_CHARS:
                            warnings.append('已达到全文处理长度上限，后续单元格未索引。')
                            break
                        add(f'{sheet.title} 第 {index + 1} 行', ' | '.join(str(cell) for cell in row if cell is not None))
            finally:
                book.close()
        else:
            from pptx import Presentation
            for index, slide in enumerate(Presentation(path).slides):
                add(f'幻灯片 {index + 1}', '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame), index + 1)
    else:
        raise ValueError('暂不支持此格式的持久化检索，请转换为 PDF 或文本。')
    if not units:
        warnings.append('没有提取到可检索的正文，原文件已保存。')
    return units, list(dict.fromkeys(warnings))
